import pandas as pd
import json
import re
import os
import sys
import shutil
import zipfile
import tempfile
import jdatetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from database import get_session, ImportBatch, ImportReview, Policy, PolicyManager, InstallmentManager, CompanyManager, Invoice, Installment, Receipt, PaymentLog, Company, load_settings

def _resolve_app_dir():
    if getattr(sys, 'frozen', False):
        return getattr(sys, '_MEIPASS', os.path.dirname(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))

def _resolve_app_data_dir():
    """پوشه ثابت ذخیره‌سازی اطلاعات و تنظیمات برنامه، مستقل از محل نصب یا اجرای فایل exe."""
    target = os.path.join("C:\\", "Bime Ba Ma", "مالی ماموت - کسر از حقوق")
    try:
        os.makedirs(target, exist_ok=True)
        return target
    except Exception:
        return _resolve_app_dir()

BASE_DIR = _resolve_app_data_dir()
DATA_ROOT = os.path.join(BASE_DIR, "MammutERP_Data")
SETTINGS_FILE = os.path.join(BASE_DIR, "settings.json")

def save_settings(data):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=4)

_BIDI_CTRL = ('\u202A','\u202B','\u202C','\u202D','\u202E','\u200E','\u200F','\u2066','\u2067','\u2068','\u2069')

RECEIPT_SUB_US = "فیش‌های کارگزاری"
RECEIPT_SUB_PASARGAD = "فیش‌های پاسارگاد"
RECEIPT_SUB_LEGACY = "فیش‌های پرداخت"

def en_digits(text_str):
    t = str(text_str)
    return t.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789"))

def to_persian_digits(value):
    text = str(value)
    for e, p in zip("0123456789", "۰۱۲۳۴۵۶۷۸۹"): text = text.replace(e, p)
    return text

def format_number(value):
    try:
        val = float(value)
        return to_persian_digits(f"{val:,.0f}")
    except:
        return to_persian_digits(value)

def format_plate_for_display(plate):
    if not plate: return ""
    p = str(plate).strip()
    for m in _BIDI_CTRL:
        p = p.replace(m, '')
    if not p: return ""
    # کلمه «ایران» باید پیش از اعمال جهت LTR برعکس نوشته شود تا پس از رندر، به‌صورت صحیح «ایران» نمایش داده شود
    p = p.replace("ایران", "ناریا")
    return f"\u202D{p}\u202C"

def _normalize_hyperlink_target(target):
    if not target or not isinstance(target, str): return target
    if not target.startswith("file:///"): return target
    stripped = target[7:]
    if len(stripped) >= 3 and stripped[0] == "/" and stripped[2] == ":":
        stripped = stripped[1:]
    return stripped.replace("/", os.sep)

def _safe_print(msg):
    try: print(msg)
    except Exception: pass

def setup_directories():
    folders = ["بایگانی ورودی", "بایگانی سالانه", "Documents/Attachments", "Documents/Receipts", "Reports/Payments", "Invoices", "Templates", "data"]
    for folder in folders: os.makedirs(os.path.join(DATA_ROOT, folder), exist_ok=True)

    if not os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
            json.dump({
                "database_path": "MammutERP_Data/data/mammut.db", "period_start_day": 25, "dark_mode": True,
                "default_installments": 9, "installment_method": "mamoot", "font_family": "Vazir",
                "target_contracts": ["1404/30000/5051"]
            }, f, ensure_ascii=False, indent=4)

def _receipt_subfolder(target):
    if target == "us": return RECEIPT_SUB_US
    if target == "pasargad": return RECEIPT_SUB_PASARGAD
    return RECEIPT_SUB_LEGACY

class ExcelEngine:
    def __init__(self, installment_count=9, installment_method="mamoot", target_contracts=None):
        self.installment_count = installment_count; self.installment_method = installment_method; self.target_contracts = target_contracts or []

    def load_excel(self, file_path): return pd.read_excel(file_path)

    def normalize_header(self, value):
        value = str(value)
        for a, b in {"\u200c": "", "‌": "", " ": "", "ي": "ی", "ك": "ک"}.items(): value = value.replace(a, b)
        return value.strip()

    def detect_policy_type(self, policy_number):
        val = str(policy_number).strip().replace(" ", "")
        if val.startswith("1"): return "ثالث"
        if val.startswith("2"): return "بدنه"
        return "نامشخص"

    def normalize_plate(self, plate):
        val = str(plate).strip().replace("(", "").replace(")", "")
        for m in _BIDI_CTRL: val = val.replace(m, '')
        if "-" in val:
            parts = [x.strip() for x in val.split("-")]
            if len(parts) == 2:
                car = parts[0].replace(" ", "")
                country = parts[1].replace("ایران", "").replace("ناریا", "").strip()
                m = re.match(r'(\d+)([^\d]+)(\d+)', car)
                if m:
                    g1, g2, g3 = m.groups()
                    p2 = g1 if len(g1) == 2 else (g3 if len(g3) == 2 else g1)
                    p3 = g3 if len(g3) == 3 else (g1 if len(g1) == 3 else g3)
                    return format_plate_for_display(f"{p2} {g2} {p3} - {country}ایران")
        return format_plate_for_display(val)

    def calculate_period(self, issue_date, cut_off_day=25):
        """
        محاسبه هوشمند دوره مالی بر اساس روز مالی انتخابی در تنظیمات:
        از 26م ماه قبل تا 25م این ماه متعلق به دوره این ماه است.
        """
        if not issue_date: return ""
        try:
            months = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
            clean_date = en_digits(str(issue_date)).split()[0].replace("-", "/")
            parts = clean_date.split("/")
            if len(parts) != 3: return ""
            y, m, d = map(int, parts)
            
            # اگر روز از روز تعیین شده (مثلا ۲۵) بزرگتر باشد به ماه بعد تعلق دارد
            if d > int(cut_off_day):
                m += 1
                if m > 12:
                    m = 1
                    y += 1
            return f"{months[m-1]} {y}"
        except Exception:
            return ""

    def detect_columns(self, raw_columns):
        mapping = {}
        aliases = {
            "insured_type": ["نوعبیمهگذار", "نوعمشتری", "نوعبیمهگذار:حقیقی-حقوقی"], 
            "contract_number": ["شمارهقرارداد", "قرارداد"],
            "policy_number": ["شمارهبیمهنامه", "شمارهبیمه"], 
            "premium_amount": ["حقبیمه", "مبلغحقبیمه", "حقبیمهبهعدد"],
            "insured_name": ["بیمهگذار", "نامبیمهگذار", "نام"], 
            "national_code": ["کدملیبیمهگذار", "کدملی"],
            "personnel_code": ["کدپرسنلی", "شمارهپرسنلی", "کدپرسنل", "شمارهپرسنل", "پرسنلی", "کداستخدامی"],
            "issue_date": ["تاریخصدوربیمهنامه", "تاریخصدور"], 
            "contract_company": ["طرفقرارداد", "شرکت"], 
            "plate": ["شمارهانتظامی", "شمارهپلاک", "پلاک"]
        }
        cols = [self.normalize_header(c) for c in raw_columns]
        for key, vals in aliases.items(): 
            mapping[key] = next((raw_columns[cols.index(v)] for v in vals if v in cols), None)
        return mapping

    @staticmethod
    def style_excel_sheet(worksheet):
        header_font = Font(name='Vazir', bold=True, size=11)
        header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        cell_font = Font(name='Vazir', bold=False, size=10)
        alignment = Alignment(horizontal='center', vertical='center')

        for cell in worksheet[1]:
            cell.font = header_font; cell.fill = header_fill; cell.alignment = alignment
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = cell_font; cell.alignment = alignment
        for col in worksheet.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            worksheet.column_dimensions[col_letter].width = max(max_len + 4, 14)

    @staticmethod
    def _set_plate_column_ltr(worksheet, df, column_name="شماره انتظامی"):
        if df is None or df.empty or column_name not in df.columns: return
        col_idx = list(df.columns).index(column_name) + 1
        try:
            ltr_align = Alignment(horizontal='center', vertical='center', readingOrder=1)
        except TypeError:
            ltr_align = Alignment(horizontal='center', vertical='center')
        for row_idx in range(2, len(df) + 2):
            cell = worksheet.cell(row=row_idx, column=col_idx)
            cell.alignment = ltr_align

    @staticmethod
    def _apply_hyperlink(cell, file_path):
        cell.hyperlink = "file:///" + file_path.replace("\\", "/")
        cell.font = Font(name='Vazir', size=10, color="0563C1", underline="single")
        cell.alignment = Alignment(horizontal='center', vertical='center')

    def export_excel_styled(self, df, path):
        rename_map = {
            "policy_number": "شماره بیمه نامه", "insured_name": "بیمه گذار", "contract_number": "شماره قرارداد",
            "contract_company": "شرکت طرف قرارداد", "national_code": "کد ملی", "issue_date": "تاریخ صدور",
            "premium_amount": "حق بیمه", "plate": "شماره انتظامی", "policy_type": "نوع بیمه",
            "personnel_code": "کد پرسنلی", "period_name": "دوره مالی"
        }
        df = df.rename(columns=rename_map)
        writer = pd.ExcelWriter(path, engine='openpyxl')
        df.to_excel(writer, index=False, sheet_name='Policies')
        ws = writer.sheets['Policies']
        self.style_excel_sheet(ws)
        self._set_plate_column_ltr(ws, df, "شماره انتظامی")
        writer.close()

    def preview(self, file_path, period_start_day=25):
        df = self.load_excel(file_path)
        mapping = self.detect_columns(df.columns)
        db = get_session()
        try:
            existing = {p[0] for p in db.query(Policy.policy_number).all()}
        finally:
            db.close()

        seen = set(); res = {"total": len(df), "approved": 0, "review": 0, "failed": 0, "rows": []}

        for idx, row in df.iterrows():
            row_data = row.to_dict()
            ins_type = str(row_data.get(mapping.get("insured_type"), "")).strip(); contract = str(row_data.get(mapping.get("contract_number"), "")).strip()
            is_haghighi = "حقیقی" in ins_type; is_valid_contract = any(c in contract for c in self.target_contracts)

            decision = {"score": 100 if is_haghighi and is_valid_contract else 0, "status": "approved" if is_haghighi and is_valid_contract else "failed", "reason": []}
            if not is_haghighi: decision["reason"].append("بیمه گذار حقوقی است")
            if not is_valid_contract: decision["reason"].append("قرارداد نامعتبر")
            if decision["status"] == "failed": res["failed"] += 1; continue

            try: prem = float(row_data.get(mapping.get("premium_amount"), 0))
            except: prem = 0
            pol_num = str(row_data.get(mapping.get("policy_number"), "")).strip()

            if pol_num in existing: decision["status"] = "failed"; decision["reason"].append("تکراری در سیستم"); res["failed"] += 1; continue
            if pol_num in seen: decision["status"] = "failed"; decision["reason"].append("تکراری در فایل"); res["failed"] += 1; continue
            if pol_num: seen.add(pol_num)

            # خواندن بدون ارور کد پرسنلی
            p_code_col = mapping.get("personnel_code")
            p_code_val = str(row_data.get(p_code_col, "")).strip() if p_code_col and pd.notna(row_data.get(p_code_col)) else ""

            issue_date = str(row_data.get(mapping.get("issue_date"), "")).strip()
            clean_data = {
                "policy_number": pol_num, "insured_name": str(row_data.get(mapping.get("insured_name"), "")).strip(),
                "contract_number": str(row_data.get(mapping.get("contract_number"), "")).strip(),
                "contract_company": str(row_data.get(mapping.get("contract_company"), "")).strip(),
                "national_code": str(row_data.get(mapping.get("national_code"), "")).strip(),
                "personnel_code": p_code_val,
                "issue_date": issue_date, "premium_amount": prem,
                "plate": self.normalize_plate(row_data.get(mapping.get("plate"), "")),
                "policy_type": self.detect_policy_type(pol_num),
                "period_name": self.calculate_period(issue_date, period_start_day)
            }
            res["rows"].append({"row": idx + 1, "data": row_data, "decision": decision, "clean_data": clean_data})
            if decision["status"] == "approved": res["approved"] += 1

        return res

    def commit_data(self, approved_data, filename="Imported.xlsx"):
        created = 0
        main_period = en_digits(approved_data[0]["period_name"]) if approved_data else "نامشخص"
        db = get_session()
        try:
            batch = ImportBatch(file_name=filename, period_name=main_period, status="processing")
            db.add(batch); db.commit(); db.refresh(batch)
            batch_id = batch.id
        finally:
            db.close()

        export_data = []
        for clean in approved_data:
            comp = CompanyManager.get_or_create(clean["contract_company"] or "نامشخص")
            pol = PolicyManager.create(comp.id, clean["policy_number"], clean["insured_name"], clean["premium_amount"], clean["contract_number"], clean["policy_type"], clean["plate"], clean["national_code"], clean["issue_date"], clean["contract_company"], clean["personnel_code"], en_digits(clean["period_name"]), batch_id)
            InstallmentManager.create(pol.id, clean["premium_amount"], self.installment_count, self.installment_method, clean["issue_date"])
            created += 1; export_data.append(clean)

        if export_data:
            df_export = pd.DataFrame(export_data)
            today_shamsi = jdatetime.date.today().strftime('%Y.%m.%d')
            input_dir = os.path.join(DATA_ROOT, f"بایگانی ورودی/Batch_{batch_id} - {today_shamsi}")
            os.makedirs(input_dir, exist_ok=True)
            self.export_excel_styled(df_export, os.path.join(input_dir, f"Batch_{batch_id}_Processed.xlsx"))

            year = en_digits(main_period.split()[1] if len(main_period.split()) > 1 else "نامشخص")
            for comp in df_export['contract_company'].unique():
                if not comp: continue
                comp_safe = "".join(c for c in comp if c.isalnum() or c in (' ', '-', '_')).strip()
                yearly_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{main_period}/{comp_safe}")
                os.makedirs(yearly_dir, exist_ok=True)
                df_comp = df_export[df_export['contract_company'] == comp]
                self.export_excel_styled(df_comp, os.path.join(yearly_dir, f"Policies_{comp_safe}.xlsx"))

        sync_error = None
        try:
            ExcelEngine.sync_master_excel()
        except Exception as e:
            sync_error = str(e)

        return {"created": created, "batch_id": batch_id, "sync_error": sync_error}

    @staticmethod
    def save_receipt_files(files_list, company_name=None, period_name=None, target=None):
        saved = []
        if company_name and company_name != "مشتریان" and period_name:
            year = en_digits(period_name.split()[1] if len(period_name.split()) > 1 else "نامشخص")
            comp_safe = "".join(c for c in company_name if c.isalnum() or c in (' ', '-', '_')).strip()
            sub = _receipt_subfolder(target)
            dest_folder = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{period_name}/{comp_safe}/{sub}")
        else:
            dest_folder = os.path.join(DATA_ROOT, "Documents/Receipts")
        os.makedirs(dest_folder, exist_ok=True)

        today_str = jdatetime.date.today().strftime('%Y-%m-%d')
        existing_nums = []
        try:
            for f in os.listdir(dest_folder):
                m = re.match(r'^(\d+)\s*-', f)
                if m:
                    try: existing_nums.append(int(m.group(1)))
                    except Exception: pass
        except Exception: pass

        next_seq = (max(existing_nums) if existing_nums else 0) + 1
        for i, f in enumerate(files_list):
            try:
                ext = os.path.splitext(f)[1]
                new_name = f"{next_seq + i} - {today_str}{ext}"
                dest = os.path.join(dest_folder, new_name)
                shutil.copy(f, dest)
                saved.append(new_name)
            except Exception: pass
        return ",".join(saved)

    @staticmethod
    def _receipt_details(db, receipt_id, pol):
        empty = {"desc": "-", "method": "-", "file_path": None, "file_name": "-"}
        if not receipt_id: return empty
        rec = db.query(Receipt).filter(Receipt.id == receipt_id).first()
        if not rec: return empty

        file_name, file_path = "-", None
        if rec.file_path:
            first_name = rec.file_path.split(",")[0].strip()
            if first_name and first_name != "-":
                file_name = first_name
                period_name = pol.period_name or "نامشخص"
                year = en_digits(period_name.split()[1] if len(period_name.split()) > 1 else "نامشخص")
                comp_safe = "".join(c for c in (pol.contract_company or "نامشخص") if c.isalnum() or c in (' ', '-', '_')).strip()
                base_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{period_name}/{comp_safe}")

                sub = _receipt_subfolder(rec.target)
                candidates = [
                    os.path.join(base_dir, sub, first_name),
                    os.path.join(base_dir, RECEIPT_SUB_LEGACY, first_name),
                    os.path.join(DATA_ROOT, "Documents", "Receipts", first_name),
                ]
                seen_paths = set()
                for c in candidates:
                    if c in seen_paths: continue
                    seen_paths.add(c)
                    if os.path.exists(c):
                        file_path = os.path.abspath(c)
                        break
        return {"desc": rec.description or "-", "method": rec.receipt_type or "-", "file_path": file_path, "file_name": file_name}

    @staticmethod
    def _attach_hyperlinks(worksheet, df, column_name, path_list):
        if df.empty or column_name not in df.columns: return
        col_idx = list(df.columns).index(column_name) + 1
        df_row_count = len(df)
        for i, path in enumerate(path_list):
            if i >= df_row_count: break
            if path and os.path.exists(path):
                cell = worksheet.cell(row=i + 2, column=col_idx)
                ExcelEngine._apply_hyperlink(cell, path)

    @staticmethod
    def _resolve_receipt_path(pol, files_str, target=None):
        if not files_str: return None
        first_file = files_str.split(",")[0].strip()
        if not first_file or first_file == "-": return None
        per = pol.period_name or "نامشخص"
        year = en_digits(per.split()[1] if len(per.split()) > 1 else "نامشخص")
        comp_safe = "".join(c for c in (pol.contract_company or "نامشخص") if c.isalnum() or c in (' ', '-', '_')).strip()
        base_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{per}/{comp_safe}")

        sub = _receipt_subfolder(target)
        candidates = [
            os.path.join(base_dir, sub, first_file),
            os.path.join(base_dir, RECEIPT_SUB_LEGACY, first_file),
            os.path.join(DATA_ROOT, "Documents", "Receipts", first_file),
        ]
        seen_paths = set()
        for c in candidates:
            if c in seen_paths: continue
            seen_paths.add(c)
            if os.path.exists(c):
                return os.path.abspath(c)
        return None

    @staticmethod
    def _get_group_due_date(period_name, inst_number):
        months = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
        try:
            m_name, y_str = en_digits(str(period_name)).split()
            m_idx = months.index(m_name) + 1
            y = int(y_str)
            m_idx += int(inst_number)
            while m_idx > 12:
                m_idx -= 12; y += 1
            return f"{y:04d}/{m_idx:02d}/15"
        except Exception:
            return ""

    @staticmethod
    def _check_file_unlocked(path):
        if not os.path.exists(path): return
        try:
            with open(path, 'a'): pass
        except PermissionError:
            raise PermissionError(f"⚠️ فایل «{os.path.basename(path)}» در حال حاضر باز است. لطفاً آن را ببندید.")

    @staticmethod
    def _atomic_write_excel(target_path, write_callback):
        ExcelEngine._check_file_unlocked(target_path)
        folder = os.path.dirname(target_path) or "."
        base = os.path.basename(target_path)
        name, ext = os.path.splitext(base)
        if not ext: ext = ".xlsx"
        os.makedirs(folder, exist_ok=True)
        tmp_path = os.path.join(folder, f"~tmp_{name}{ext}")

        if os.path.exists(tmp_path):
            try: os.remove(tmp_path)
            except Exception: pass

        try:
            write_callback(tmp_path)
        except Exception as e:
            if os.path.exists(tmp_path):
                try: os.remove(tmp_path)
                except Exception: pass
            raise Exception(f"خطا در نوشتن فایل موقت: {e}")

        try:
            os.replace(tmp_path, target_path)
        except Exception as e:
            try:
                if os.path.exists(target_path): os.remove(target_path)
                os.rename(tmp_path, target_path)
            except Exception as e2:
                if os.path.exists(tmp_path):
                    try: os.remove(tmp_path)
                    except Exception: pass
                raise Exception(f"خطا در جایگزینی فایل: {e} | {e2}")

    @staticmethod
    def sync_master_excel():
        os.makedirs(DATA_ROOT, exist_ok=True)
        master_path = os.path.join(DATA_ROOT, "بانک جامع اطلاعات و اقساط کل.xlsx")
        data = InstallmentManager.get_all_with_policies()

        policies_dict = {}
        for inst, pol in data:
            if pol.id not in policies_dict:
                policies_dict[pol.id] = {
                    "ردیف": len(policies_dict) + 1,
                    "شماره بیمه نامه": pol.policy_number,
                    "بیمه گذار": pol.insured_name,
                    "شرکت طرف قرارداد": pol.contract_company or "نامشخص",
                    "دوره مالی": pol.period_name or "نامشخص",
                    "حق بیمه": pol.premium_amount,
                    "شماره انتظامی": format_plate_for_display(pol.plate),
                    "کد ملی": pol.national_code,
                    "کد پرسنلی": pol.personnel_code or "-",
                    "تاریخ صدور": pol.issue_date,
                    "وضعیت صورتحساب": "دارای صورتحساب" if pol.is_invoiced else "بدون صورتحساب",
                    "شماره صورتحساب": pol.invoice_number or "-",
                }
        df_policies = pd.DataFrame(list(policies_dict.values())) if policies_dict else pd.DataFrame()

        db = get_session()
        inst_records = []
        us_file_paths, pas_file_paths = [], []
        payment_history_records = []
        try:
            for idx, (inst, pol) in enumerate(data):
                paid_us = inst.paid_amount_us or 0.0
                paid_pas = inst.paid_amount_pasargad or 0.0
                rem_us = max(0.0, inst.amount - paid_us)
                rem_pas = max(0.0, inst.amount - paid_pas)

                stat_us = "پرداخت شده" if paid_us >= inst.amount - 0.5 else ("مانده دار" if paid_us > 0 else "پرداخت نشده")
                stat_pas = "پرداخت شده" if paid_pas >= inst.amount - 0.5 else ("مانده دار" if paid_pas > 0 else "پرداخت نشده")
                stat_final = "تسویه کامل و نهایی" if paid_pas >= inst.amount - 0.5 else "جاری"

                us_info = ExcelEngine._receipt_details(db, inst.receipt_us_id, pol)
                pas_info = ExcelEngine._receipt_details(db, inst.receipt_pasargad_id, pol)
                us_file_paths.append(us_info["file_path"])
                pas_file_paths.append(pas_info["file_path"])

                inst_records.append({
                    "ردیف": idx + 1, "شناسه ۱۰ رقمی": inst.tracking_id,
                    "شماره بیمه نامه": pol.policy_number, "بیمه گذار": pol.insured_name,
                    "شرکت طرف قرارداد": pol.contract_company or "نامشخص",
                    "دوره مالی": pol.period_name or "نامشخص",
                    "شماره قسط": inst.installment_number, "سررسید": inst.due_date,
                    "مبلغ کل قسط": inst.amount,
                    "پرداختی به کارگزاری": paid_us, "مانده کارگزاری": rem_us, "وضعیت کارگزاری": stat_us,
                    "روش پرداخت کارگزاری": us_info["method"], "توضیحات کارگزاری": us_info["desc"], "فیش کارگزاری": us_info["file_name"],
                    "پرداختی به پاسارگاد": paid_pas, "مانده پاسارگاد": rem_pas, "وضعیت پاسارگاد": stat_pas,
                    "روش پرداخت پاسارگاد": pas_info["method"], "توضیحات پاسارگاد": pas_info["desc"], "فیش پاسارگاد": pas_info["file_name"],
                    "وضعیت نهایی": stat_final,
                })

            try:
                logs = db.query(PaymentLog).order_by(PaymentLog.paid_at.desc()).all()
                if logs:
                    inst_map = {inst.id: (inst, pol) for inst, pol in data}
                    for lg_idx, lg in enumerate(logs, start=1):
                        inst_pol = inst_map.get(lg.installment_id)
                        if not inst_pol: continue
                        lg_inst, lg_pol = inst_pol
                        target_txt = "واریز به کارگزاری" if lg.target == "us" else "پرداخت به پاسارگاد"
                        try: paid_at_txt = lg.paid_at.strftime('%Y/%m/%d %H:%M:%S') if lg.paid_at else "-"
                        except Exception: paid_at_txt = str(lg.paid_at) if lg.paid_at else "-"
                        payment_history_records.append({
                            "ردیف": lg_idx,
                            "تاریخ و ساعت ثبت": to_persian_digits(paid_at_txt),
                            "کاربر ثبت‌کننده": lg.registered_by or "کاربر سیستم",
                            "مقصد": target_txt,
                            "مبلغ پرداختی": lg.amount or 0.0,
                            "شناسه ۱۰ رقمی قسط": lg_inst.tracking_id,
                            "شماره بیمه نامه": lg_pol.policy_number,
                            "بیمه گذار": lg_pol.insured_name,
                            "شرکت طرف قرارداد": lg_pol.contract_company or "نامشخص",
                            "دوره مالی": lg_pol.period_name or "نامشخص",
                            "شماره قسط": lg_inst.installment_number,
                            "مبلغ کل قسط": lg_inst.amount,
                            "شناسه رسید": lg.receipt_id or "-",
                        })
            except Exception: pass
        finally:
            db.close()

        df_installments = pd.DataFrame(inst_records) if inst_records else pd.DataFrame()
        df_payment_history = pd.DataFrame(payment_history_records) if payment_history_records else pd.DataFrame()

        company_groups = {}
        for inst, pol in data:
            comp = (pol.contract_company or "نامشخص").strip()
            per = en_digits(pol.period_name or "نامشخص").strip()
            num = inst.installment_number
            key = (comp, per, num)
            if key not in company_groups:
                company_groups[key] = {"count": 0, "amount": 0.0, "paid_us": 0.0, "paid_pas": 0.0, "invoiced_all": True}
            g = company_groups[key]
            g["count"] += 1
            g["amount"] += inst.amount or 0.0
            g["paid_us"] += (inst.paid_amount_us or 0.0)
            g["paid_pas"] += (inst.paid_amount_pasargad or 0.0)
            if not pol.is_invoiced: g["invoiced_all"] = False

        company_records = []
        for idx, ((comp, per, num), g) in enumerate(sorted(company_groups.items(), key=lambda x: (x[0][0], x[0][1], x[0][2])), start=1):
            rem_us = max(0.0, g["amount"] - g["paid_us"])
            rem_pas = max(0.0, g["amount"] - g["paid_pas"])
            stat_us = "پرداخت شده" if g["paid_us"] >= g["amount"] - 0.5 else ("مانده دار" if g["paid_us"] > 0 else "پرداخت نشده")
            stat_pas = "پرداخت شده" if g["paid_pas"] >= g["amount"] - 0.5 else ("مانده دار" if g["paid_pas"] > 0 else "پرداخت نشده")
            stat_final = "تسویه کامل و نهایی" if g["paid_pas"] >= g["amount"] - 0.5 else "جاری"
            company_records.append({
                "ردیف": idx, "شرکت/سازمان": comp, "دوره مالی": per, "شماره قسط": num,
                "تعداد بیمه": g["count"], "سررسید شرکتی": ExcelEngine._get_group_due_date(per, num),
                "مبلغ کل قسط": g["amount"],
                "پرداختی به کارگزاری": g["paid_us"], "مانده کارگزاری": rem_us, "وضعیت کارگزاری": stat_us,
                "پرداختی به پاسارگاد": g["paid_pas"], "مانده پاسارگاد": rem_pas, "وضعیت پاسارگاد": stat_pas,
                "وضعیت نهایی": stat_final,
                "وضعیت صورتحساب": "دارای صورتحساب" if g["invoiced_all"] else "بدون صورتحساب",
            })
        df_company = pd.DataFrame(company_records) if company_records else pd.DataFrame()

        log_folder = os.path.join(DATA_ROOT, "Reports", "Payments")
        all_logs = []
        if os.path.exists(log_folder):
            for f in os.listdir(log_folder):
                if f.endswith(".xlsx") and not f.startswith("~tmp_"):
                    try:
                        df_l = pd.read_excel(os.path.join(log_folder, f))
                        if not df_l.empty: all_logs.append(df_l)
                    except Exception: pass
        df_all_logs = pd.concat(all_logs, ignore_index=True, sort=False) if all_logs else None

        def _write(out_path):
            with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
                wrote_any = False
                if not df_policies.empty:
                    df_policies.to_excel(writer, index=False, sheet_name='بیمه‌نامه‌ها')
                    ws = writer.sheets['بیمه‌نامه‌ها']
                    ExcelEngine.style_excel_sheet(ws)
                    ExcelEngine._set_plate_column_ltr(ws, df_policies, "شماره انتظامی")
                    wrote_any = True
                if not df_installments.empty:
                    df_installments.to_excel(writer, index=False, sheet_name='مالی اشخاص')
                    ws_inst = writer.sheets['مالی اشخاص']
                    ExcelEngine.style_excel_sheet(ws_inst)
                    ExcelEngine._attach_hyperlinks(ws_inst, df_installments, "فیش کارگزاری", us_file_paths)
                    ExcelEngine._attach_hyperlinks(ws_inst, df_installments, "فیش پاسارگاد", pas_file_paths)
                    wrote_any = True
                if not df_company.empty:
                    df_company.to_excel(writer, index=False, sheet_name='مالی شرکت')
                    ExcelEngine.style_excel_sheet(writer.sheets['مالی شرکت'])
                    wrote_any = True
                if df_all_logs is not None and not df_all_logs.empty:
                    df_all_logs.to_excel(writer, index=False, sheet_name='ریز تراکنش‌های پرداختی')
                    ExcelEngine.style_excel_sheet(writer.sheets['ریز تراکنش‌های پرداختی'])
                    wrote_any = True
                if not df_payment_history.empty:
                    df_payment_history.to_excel(writer, index=False, sheet_name='ریز پرداخت‌ها')
                    ExcelEngine.style_excel_sheet(writer.sheets['ریز پرداخت‌ها'])
                    wrote_any = True
                if not wrote_any:
                    pd.DataFrame({"پیام": ["هیچ داده‌ای وجود ندارد"]}).to_excel(writer, index=False, sheet_name='خالی')

        ExcelEngine._atomic_write_excel(master_path, _write)

    @staticmethod
    def sync_company_excel(company_name):
        db = get_session()
        try:
            rows = db.query(Installment, Policy).join(Policy).filter(Policy.contract_company == company_name).all()
            if not rows: return

            by_period = {}
            for inst, pol in rows:
                by_period.setdefault(pol.period_name or "نامشخص", []).append((inst, pol))

            comp_safe = "".join(c for c in company_name if c.isalnum() or c in (' ', '-', '_')).strip()

            for period_name, items in by_period.items():
                year = en_digits(period_name.split()[1] if len(period_name.split()) > 1 else "نامشخص")
                comp_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{period_name}/{comp_safe}")
                os.makedirs(comp_dir, exist_ok=True)
                out_path = os.path.join(comp_dir, f"Policies_{comp_safe}.xlsx")

                policies_dict = {}
                inst_records = []
                us_file_paths, pas_file_paths = [], []
                payment_history_records = []
                for idx, (inst, pol) in enumerate(items):
                    if pol.id not in policies_dict:
                        policies_dict[pol.id] = {
                            "ردیف": len(policies_dict) + 1,
                            "شماره بیمه نامه": pol.policy_number, "بیمه گذار": pol.insured_name,
                            "دوره مالی": pol.period_name or "نامشخص", "حق بیمه": pol.premium_amount,
                            "شماره انتظامی": format_plate_for_display(pol.plate),
                            "کد ملی": pol.national_code,
                            "کد پرسنلی": pol.personnel_code or "-",
                            "تاریخ صدور": pol.issue_date,
                            "وضعیت صورتحساب": "دارای صورتحساب" if pol.is_invoiced else "بدون صورتحساب",
                            "شماره صورتحساب": pol.invoice_number or "-",
                        }
                    paid_us = inst.paid_amount_us or 0.0
                    paid_pas = inst.paid_amount_pasargad or 0.0
                    rem_us = max(0.0, inst.amount - paid_us)
                    rem_pas = max(0.0, inst.amount - paid_pas)

                    us_info = ExcelEngine._receipt_details(db, inst.receipt_us_id, pol)
                    pas_info = ExcelEngine._receipt_details(db, inst.receipt_pasargad_id, pol)
                    us_file_paths.append(us_info["file_path"])
                    pas_file_paths.append(pas_info["file_path"])

                    inst_records.append({
                        "ردیف": idx + 1, "شناسه ۱۰ رقمی": inst.tracking_id,
                        "شماره بیمه نامه": pol.policy_number, "بیمه گذار": pol.insured_name,
                        "شماره قسط": inst.installment_number, "سررسید": inst.due_date, "مبلغ کل قسط": inst.amount,
                        "پرداختی به کارگزاری": paid_us, "مانده کارگزاری": rem_us,
                        "وضعیت کارگزاری": "پرداخت شده" if paid_us >= inst.amount - 0.5 else ("مانده دار" if paid_us > 0 else "پرداخت نشده"),
                        "روش پرداخت کارگزاری": us_info["method"], "توضیحات کارگزاری": us_info["desc"], "فیش کارگزاری": us_info["file_name"],
                        "پرداختی به پاسارگاد": paid_pas, "مانده پاسارگاد": rem_pas,
                        "وضعیت پاسارگاد": "پرداخت شده" if paid_pas >= inst.amount - 0.5 else ("مانده دار" if paid_pas > 0 else "پرداخت نشده"),
                        "روش پرداخت پاسارگاد": pas_info["method"], "توضیحات پاسارگاد": pas_info["desc"], "فیش پاسارگاد": pas_info["file_name"],
                    })

                try:
                    inst_ids = [inst.id for inst, pol in items]
                    if inst_ids:
                        logs = db.query(PaymentLog).filter(PaymentLog.installment_id.in_(inst_ids)).order_by(PaymentLog.paid_at.desc()).all()
                        inst_map = {inst.id: (inst, pol) for inst, pol in items}
                        for lg_idx, lg in enumerate(logs, start=1):
                            ip = inst_map.get(lg.installment_id)
                            if not ip: continue
                            lg_inst, lg_pol = ip
                            target_txt = "واریز به کارگزاری" if lg.target == "us" else "پرداخت به پاسارگاد"
                            try: paid_at_txt = lg.paid_at.strftime('%Y/%m/%d %H:%M:%S') if lg.paid_at else "-"
                            except Exception: paid_at_txt = str(lg.paid_at) if lg.paid_at else "-"
                            payment_history_records.append({
                                "ردیف": lg_idx,
                                "تاریخ و ساعت ثبت": to_persian_digits(paid_at_txt),
                                "کاربر ثبت‌کننده": lg.registered_by or "کاربر سیستم",
                                "مقصد": target_txt,
                                "مبلغ پرداختی": lg.amount or 0.0,
                                "شناسه ۱۰ رقمی قسط": lg_inst.tracking_id,
                                "شماره بیمه نامه": lg_pol.policy_number,
                                "بیمه گذار": lg_pol.insured_name,
                                "شماره قسط": lg_inst.installment_number,
                                "مبلغ کل قسط": lg_inst.amount,
                            })
                except Exception: pass

                company_groups = {}
                for inst, pol in items:
                    per = en_digits(pol.period_name or "نامشخص").strip()
                    num = inst.installment_number
                    key = (per, num)
                    if key not in company_groups:
                        company_groups[key] = {"count": 0, "amount": 0.0, "paid_us": 0.0, "paid_pas": 0.0, "invoiced_all": True}
                    g = company_groups[key]
                    g["count"] += 1
                    g["amount"] += inst.amount or 0.0
                    g["paid_us"] += (inst.paid_amount_us or 0.0)
                    g["paid_pas"] += (inst.paid_amount_pasargad or 0.0)
                    if not pol.is_invoiced: g["invoiced_all"] = False

                company_records = []
                for idx, ((per, num), g) in enumerate(sorted(company_groups.items(), key=lambda x: (x[0][0], x[0][1])), start=1):
                    rem_us = max(0.0, g["amount"] - g["paid_us"])
                    rem_pas = max(0.0, g["amount"] - g["paid_pas"])
                    stat_us = "پرداخت شده" if g["paid_us"] >= g["amount"] - 0.5 else ("مانده دار" if g["paid_us"] > 0 else "پرداخت نشده")
                    stat_pas = "پرداخت شده" if g["paid_pas"] >= g["amount"] - 0.5 else ("مانده دار" if g["paid_pas"] > 0 else "پرداخت نشده")
                    stat_final = "تسویه کامل و نهایی" if g["paid_pas"] >= g["amount"] - 0.5 else "جاری"
                    company_records.append({
                        "ردیف": idx, "شرکت/سازمان": company_name, "دوره مالی": per, "شماره قسط": num,
                        "تعداد بیمه": g["count"], "سررسید شرکتی": ExcelEngine._get_group_due_date(per, num),
                        "مبلغ کل قسط": g["amount"],
                        "پرداختی به کارگزاری": g["paid_us"], "مانده کارگزاری": rem_us, "وضعیت کارگزاری": stat_us,
                        "پرداختی به پاسارگاد": g["paid_pas"], "مانده پاسارگاد": rem_pas, "وضعیت پاسارگاد": stat_pas,
                        "وضعیت نهایی": stat_final,
                        "وضعیت صورتحساب": "دارای صورتحساب" if g["invoiced_all"] else "بدون صورتحساب",
                    })

                df_policies = pd.DataFrame(list(policies_dict.values()))
                df_installments = pd.DataFrame(inst_records)
                df_company = pd.DataFrame(company_records) if company_records else pd.DataFrame()
                df_history = pd.DataFrame(payment_history_records) if payment_history_records else pd.DataFrame()

                def _write(out_path, _pol=df_policies, _ins=df_installments, _com=df_company, _hist=df_history,
                           _usfp=us_file_paths, _pasfp=pas_file_paths):
                    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
                        _pol.to_excel(writer, index=False, sheet_name='بیمه‌نامه‌ها')
                        ws = writer.sheets['بیمه‌نامه‌ها']
                        ExcelEngine.style_excel_sheet(ws)
                        ExcelEngine._set_plate_column_ltr(ws, _pol, "شماره انتظامی")

                        _ins.to_excel(writer, index=False, sheet_name='مالی اشخاص')
                        ws_inst = writer.sheets['مالی اشخاص']
                        ExcelEngine.style_excel_sheet(ws_inst)
                        ExcelEngine._attach_hyperlinks(ws_inst, _ins, "فیش کارگزاری", _usfp)
                        ExcelEngine._attach_hyperlinks(ws_inst, _ins, "فیش پاسارگاد", _pasfp)

                        if not _com.empty:
                            _com.to_excel(writer, index=False, sheet_name='مالی شرکت')
                            ExcelEngine.style_excel_sheet(writer.sheets['مالی شرکت'])

                        if not _hist.empty:
                            _hist.to_excel(writer, index=False, sheet_name='ریز پرداخت‌ها')
                            ExcelEngine.style_excel_sheet(writer.sheets['ریز پرداخت‌ها'])

                ExcelEngine._atomic_write_excel(out_path, _write)
        finally:
            db.close()

    @staticmethod
    def export_payment_log(company, data_to_pay, target, method, notes, files_str):
        folder = os.path.join(DATA_ROOT, "Reports", "Payments")
        os.makedirs(folder, exist_ok=True)
        today_shamsi = jdatetime.date.today().strftime('%Y%m%d')
        path = os.path.join(folder, f"Payment_{company}_{today_shamsi}.xlsx")

        username = load_settings().get("user_name", "کاربر سیستم")
        now_str = jdatetime.datetime.now().strftime('%Y/%m/%d %H:%M:%S')

        records = []
        new_file_paths = []
        for inst, pol, paid_amt in data_to_pay:
            paid_us = inst.paid_amount_us or 0.0
            paid_pas = inst.paid_amount_pasargad or 0.0
            rem_us = max(0.0, inst.amount - paid_us)
            rem_pas = max(0.0, inst.amount - paid_pas)

            first_file_path = ExcelEngine._resolve_receipt_path(pol, files_str, target=target)
            new_file_paths.append(first_file_path)

            records.append({
                "کاربر ثبت کننده": username,
                "شناسه ۱۰ رقمی قسط": inst.tracking_id,
                "شماره بیمه نامه": pol.policy_number,
                "بیمه گذار": pol.insured_name,
                "شرکت طرف قرارداد": pol.contract_company or "نامشخص",
                "دوره مالی": pol.period_name or "نامشخص",
                "شماره قسط": inst.installment_number,
                "مبلغ کل قسط": inst.amount,
                "مبلغ پرداختی این تراکنش": paid_amt,
                "پرداختی کل به کارگزاری": paid_us,
                "مانده نزد کارگزاری": rem_us,
                "وضعیت کارگزاری": "تسویه شده" if rem_us <= 0.5 else "مانده دار",
                "پرداختی کل به پاسارگاد": paid_pas,
                "مانده نزد پاسارگاد": rem_pas,
                "وضعیت پاسارگاد": "تسویه شده" if rem_pas <= 0.5 else "مانده دار",
                "مقصد این تراکنش": "واریز به کارگزاری" if target == "us" else "پرداخت به پاسارگاد",
                "روش پرداخت": method,
                "تاریخ و ساعت ثبت": now_str,
                "توضیحات": notes or "-",
                "فایل ضمیمه": files_str or "-"
            })
        df_new = pd.DataFrame(records)

        existing_df = None
        existing_hyperlinks = []
        if os.path.exists(path):
            try:
                wb = openpyxl.load_workbook(path)
                if 'Payments' in wb.sheetnames:
                    ws = wb['Payments']
                    headers = [c.value for c in ws[1]]
                    file_col_idx = headers.index("فایل ضمیمه") + 1 if "فایل ضمیمه" in headers else None

                    rows_data = []
                    for r in range(2, ws.max_row + 1):
                        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
                        rows_data.append(row_vals)
                        if file_col_idx:
                            cell = ws.cell(row=r, column=file_col_idx)
                            existing_hyperlinks.append(
                                _normalize_hyperlink_target(cell.hyperlink.target) if cell.hyperlink else None
                            )
                        else:
                            existing_hyperlinks.append(None)
                    if rows_data:
                        existing_df = pd.DataFrame(rows_data, columns=headers)
                wb.close()
            except Exception:
                existing_df = None
                existing_hyperlinks = []

        if existing_df is not None:
            for col in df_new.columns:
                if col not in existing_df.columns:
                    existing_df[col] = None
            existing_df = existing_df[df_new.columns]
            df_all = pd.concat([existing_df, df_new], ignore_index=True)
            all_file_paths = existing_hyperlinks + new_file_paths
        else:
            df_all = df_new
            all_file_paths = new_file_paths

        def _write(out_path):
            with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
                df_all.to_excel(writer, index=False, sheet_name='Payments')
                ws = writer.sheets['Payments']
                ExcelEngine.style_excel_sheet(ws)

                if "فایل ضمیمه" in df_all.columns:
                    file_col_idx = list(df_all.columns).index("فایل ضمیمه") + 1
                    df_row_count = len(df_all)
                    for i, fp in enumerate(all_file_paths):
                        if i >= df_row_count: break
                        if fp and os.path.exists(fp):
                            cell = ws.cell(row=i + 2, column=file_col_idx)
                            ExcelEngine._apply_hyperlink(cell, fp)

        sync_errors = []
        try:
            ExcelEngine._atomic_write_excel(path, _write)
        except Exception as e:
            sync_errors.append(f"ثبت تراکنش ({os.path.basename(path)}): {e}")

        real_companies = {(pol.contract_company or "").strip() for _, pol, _ in data_to_pay if pol.contract_company}
        if not real_companies and company:
            real_companies = {company.strip()}
        for comp_name in real_companies:
            try:
                ExcelEngine.sync_company_excel(comp_name)
            except Exception as e:
                sync_errors.append(f"شرکت {comp_name}: {e}")

        try:
            ExcelEngine.sync_master_excel()
        except Exception as e:
            sync_errors.append(f"جامع: {e}")

        return {"success": len(sync_errors) == 0, "error": "; ".join(sync_errors) if sync_errors else None}

    # ==================== تولید صورتحساب هوشمند و پایدار ====================
    @staticmethod
    def generate_invoice(inv_type, template_path, context_data, period_name, company_name):
        try:
            from docxtpl import DocxTemplate
            import docx
            from docx.shared import Pt
            from docx.oxml.ns import qn, nsdecls
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.table import WD_TABLE_ALIGNMENT
            from docx.oxml import OxmlElement, parse_xml
        except ImportError:
            raise Exception("لطفاً کتابخانه‌های مورد نیاز را با دستور زیر نصب کنید:\npip install docxtpl python-docx comtypes")

        doc = DocxTemplate(template_path)
        
        placeholder_marker = "__INVOICE_TABLE_PLACEHOLDER__"
        context_data["invoice_table"] = placeholder_marker
        doc.render(context_data)

        year = en_digits(period_name.split()[1] if len(period_name.split()) > 1 else "نامشخص")
        comp_safe = "".join(c for c in company_name if c.isalnum() or c in (' ', '-', '_')).strip()

        # ذخیره‌سازی نوع ۱ و ۳ در پوشه دوره، نوع ۲ در پوشه همان شرکت
        if inv_type == 2:
            save_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{period_name}/{comp_safe}")
        else:
            save_dir = os.path.join(DATA_ROOT, f"بایگانی سالانه/{year}/{period_name}")
        os.makedirs(save_dir, exist_ok=True)

        today_shamsi = jdatetime.date.today().strftime('%Y.%m.%d')
        base_filename = f"صورتحساب - {period_name} - {today_shamsi}"
        file_name = base_filename
        counter = 1
        while os.path.exists(os.path.join(save_dir, f"{file_name}.docx")):
            file_name = f"{base_filename} ({counter})"
            counter += 1

        temp_docx = os.path.join(save_dir, f"~tmp_{file_name}.docx")
        doc.save(temp_docx)

        real_doc = docx.Document(temp_docx)
        target_p = None
        for p in real_doc.paragraphs:
            if placeholder_marker in p.text:
                target_p = p
                p.text = p.text.replace(placeholder_marker, "")
                break

        if target_p is None:
            for tbl in real_doc.tables:
                for row in tbl.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            if placeholder_marker in p.text:
                                target_p = p
                                p.text = p.text.replace(placeholder_marker, "")
                                break
                        if target_p: break
                    if target_p: break
                if target_p: break

        if target_p is not None:
            if inv_type == 1:
                headers = ["ردیف", "شرکت طرف قرارداد", "تعداد ثالث", "تعداد بدنه", "مبلغ کل حق بیمه", "تاریخ صدور"]
            elif inv_type == 2:
                headers = ["ردیف", "نوع بیمه", "نام پرسنل", "کد پرسنلی", "کد ملی", "مبلغ حق بیمه", "تاریخ صدور"]
            else: # نوع ۳: تجمیعی حقیقی
                headers = ["ردیف", "نام پرسنل", "نام شرکت", "کد پرسنلی", "کد ملی", "شماره بیمه نامه", "تاریخ صدور", "حق بیمه (ریال)"]

            table = real_doc.add_table(rows=1, cols=len(headers))
            
            # رفع ارور Table Grid با تزریق کادر مستقیم XML
            try:
                borders = parse_xml(r'''
                    <w:tblBorders %s>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>
                    </w:tblBorders>
                ''' % nsdecls('w'))
                table._tbl.tblPr.append(borders)
            except Exception:
                pass

            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            try:
                bidi = OxmlElement('w:bidiVisual')
                table._tbl.tblPr.append(bidi)
            except Exception:
                pass

            def format_docx_cell(cell, text_val, is_bold=False, font_size=8):
                cell.text = str(text_val)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Vazir'
                    r.font.size = Pt(font_size)
                    r.font.bold = is_bold
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Vazir')
                    r._element.rPr.rFonts.set(qn('w:cs'), 'Vazir')
                try:
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcPr.append(parse_xml(r'<w:noWrap %s/>' % nsdecls('w')))
                except Exception:
                    pass

            # هدر جدول با سایز 9pt
            for i, h in enumerate(headers):
                format_docx_cell(table.rows[0].cells[i], h, is_bold=True, font_size=9)

            rows_items = []
            for idx, item in enumerate(context_data["items"], start=1):
                row_cells = table.add_row().cells
                try:
                    trPr = table.rows[-1]._tr.get_or_add_trPr()
                    trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
                except Exception:
                    pass

                if inv_type == 1:
                    vals = [
                        to_persian_digits(idx),
                        str(item.get('company', '')),
                        str(item.get('third_count', '۰')),
                        str(item.get('body_count', '۰')),
                        str(item.get('premium', '۰')),
                        to_persian_digits(str(item.get('date', '')))
                    ]
                elif inv_type == 2:
                    vals = [
                        to_persian_digits(idx),
                        str(item.get('type', '')),
                        str(item.get('name', '')),
                        to_persian_digits(str(item.get('personnel_code', ''))),
                        to_persian_digits(str(item.get('national_code', ''))),
                        str(item.get('premium', '۰')),
                        to_persian_digits(str(item.get('date', '')))
                    ]
                else: # نوع ۳: تجمیعی حقیقی
                    vals = [
                        to_persian_digits(idx),
                        str(item.get('name', '')),
                        str(item.get('company', '')),
                        to_persian_digits(str(item.get('personnel_code', ''))),
                        to_persian_digits(str(item.get('national_code', ''))),
                        to_persian_digits(str(item.get('policy_number', ''))),
                        to_persian_digits(str(item.get('date', ''))),
                        str(item.get('premium', '۰'))
                    ]

                rows_items.append(vals)
                # داده‌های سطرها با سایز 8pt
                for col_i, val in enumerate(vals):
                    format_docx_cell(row_cells[col_i], val, is_bold=False, font_size=8)

            # افزایش عرض جدول تا 600pt جهت جلوگیری از دوخطی شدن ستون‌ها
            total_table_width_pt = 600
            max_lens = [len(str(h)) for h in headers]
            for r_val in rows_items:
                for i, v in enumerate(r_val):
                    max_lens[i] = max(max_lens[i], len(str(v)))

            sum_lens = sum(max_lens) if sum(max_lens) > 0 else 1
            calculated_widths = [max(int((l / sum_lens) * total_table_width_pt), 32) for l in max_lens]
            w_sum = sum(calculated_widths)
            col_widths = [int(w * total_table_width_pt / w_sum) for w in calculated_widths]

            for row in table.rows:
                for idx_col, cell in enumerate(row.cells):
                    cell.width = Pt(col_widths[idx_col])

            target_p._p.addnext(table._tbl)

        docx_out_path = os.path.join(save_dir, f"{file_name}.docx")
        pdf_out_path = os.path.join(save_dir, f"{file_name}.pdf")
        real_doc.save(docx_out_path)

        if os.path.exists(temp_docx):
            try: os.remove(temp_docx)
            except Exception: pass

        # تبدیل امن و پایدار به PDF با پشتیبانی چندگانه جهت خروجی EXE
        pdf_success = False
        # 1. تلاش با win32com
        try:
            import win32com.client
            import pythoncom
            pythoncom.CoInitialize()
            word_app = win32com.client.DispatchEx("Word.Application")
            word_app.Visible = False
            word_app.DisplayAlerts = False
            doc_obj = word_app.Documents.Open(os.path.abspath(docx_out_path))
            doc_obj.SaveAs(os.path.abspath(pdf_out_path), FileFormat=17)
            doc_obj.Close(False)
            word_app.Quit()
            pythoncom.CoUninitialize()
            pdf_success = True
        except Exception:
            # 2. تلاش با comtypes
            try:
                import comtypes
                import comtypes.client
                comtypes.CoInitialize()
                word_app = comtypes.client.CreateObject('Word.Application')
                word_app.Visible = False
                doc_obj = word_app.Documents.Open(os.path.abspath(docx_out_path))
                doc_obj.SaveAs(os.path.abspath(pdf_out_path), FileFormat=17)
                doc_obj.Close(False)
                word_app.Quit()
                comtypes.CoUninitialize()
                pdf_success = True
            except Exception:
                # 3. تلاش با docx2pdf
                try:
                    from docx2pdf import convert
                    convert(os.path.abspath(docx_out_path), os.path.abspath(pdf_out_path))
                    pdf_success = True
                except Exception as e3:
                    _safe_print(f"[PDF Conversion Warning] {e3}")

        return docx_out_path, (pdf_out_path if pdf_success else None), save_dir

class BackupManager:
    """
    مدیریت خروجی گرفتن (Export) و ورود اطلاعات (Import) کامل یا بخشی از سیستم.
    خروجی شامل یک فایل ZIP است که هم رکوردهای دیتابیس (به صورت JSON) و هم
    تمامی فایل‌های بایگانی، اکسل‌ها، صورتحساب‌ها و رسیدهای مرتبط را در بر می‌گیرد.
    """

    ARCHIVE_FOLDER = "بایگانی سالانه"
    ALWAYS_INCLUDE_FOLDERS = ["Invoices", "Templates", "Reports"]

    @staticmethod
    def _report(progress_cb, pct, msg):
        if progress_cb:
            try: progress_cb(int(pct), msg)
            except Exception: pass

    @staticmethod
    def get_available_periods():
        months_order = {"فروردین": 1, "اردیبهشت": 2, "خرداد": 3, "تیر": 4, "مرداد": 5, "شهریور": 6,
                         "مهر": 7, "آبان": 8, "آذر": 9, "دی": 10, "بهمن": 11, "اسفند": 12}
        db = get_session()
        try:
            rows = db.query(Policy.period_name).distinct().all()
            periods = sorted({r[0] for r in rows if r[0]}, key=lambda p: BackupManager._period_sort_key(p, months_order))
            return periods
        finally:
            db.close()

    @staticmethod
    def _period_sort_key(period_name, months_order):
        try:
            parts = str(period_name).split()
            if len(parts) >= 2:
                return (int(en_digits(parts[1])), months_order.get(parts[0], 0))
        except Exception:
            pass
        return (0, 0)

    @staticmethod
    def _period_year(period_name):
        try:
            parts = str(period_name).split()
            if len(parts) >= 2:
                return en_digits(parts[1])
        except Exception:
            pass
        return None

    # ------------------------- سریالایز کردن رکوردها -------------------------
    @staticmethod
    def _ser_company(c):
        return {"id": c.id, "name": c.name, "code": c.code, "phone": c.phone}

    @staticmethod
    def _ser_policy(p):
        return {
            "id": p.id, "company_id": p.company_id, "source_batch_id": p.source_batch_id,
            "policy_number": p.policy_number, "contract_number": p.contract_number,
            "insured_name": p.insured_name, "premium_amount": p.premium_amount,
            "issue_date": p.issue_date, "national_code": p.national_code, "plate": p.plate,
            "personnel_code": p.personnel_code, "policy_type": p.policy_type,
            "contract_company": p.contract_company, "period_name": p.period_name,
            "is_invoiced": bool(p.is_invoiced), "invoice_number": p.invoice_number,
        }

    @staticmethod
    def _ser_installment(i):
        return {
            "id": i.id, "tracking_id": i.tracking_id, "policy_id": i.policy_id,
            "installment_number": i.installment_number, "amount": i.amount,
            "paid_amount_us": i.paid_amount_us, "paid_amount_pasargad": i.paid_amount_pasargad,
            "due_date": i.due_date, "status_us": i.status_us, "status_pasargad": i.status_pasargad,
            "receipt_us_id": i.receipt_us_id, "receipt_pasargad_id": i.receipt_pasargad_id,
        }

    @staticmethod
    def _ser_receipt(r):
        return {
            "id": r.id, "receipt_type": r.receipt_type, "amount": r.amount,
            "payment_date": r.payment_date, "description": r.description,
            "file_path": r.file_path, "target": r.target,
        }

    @staticmethod
    def _ser_payment_log(lg):
        return {
            "installment_id": lg.installment_id, "receipt_id": lg.receipt_id,
            "target": lg.target, "amount": lg.amount, "registered_by": lg.registered_by,
        }

    @staticmethod
    def _ser_invoice(inv):
        return {
            "invoice_number": inv.invoice_number, "company_name": inv.company_name,
            "period_name": inv.period_name, "total_policies": inv.total_policies,
            "total_amount": inv.total_amount, "created_date": inv.created_date,
            "folder_path": inv.folder_path, "status": inv.status,
        }

    # ------------------------------- Export -------------------------------
    @staticmethod
    def export_backup(scope, params, dest_zip_path, progress_cb=None):
        """
        scope: "all" | "range" | "period"
        params: {"date_from":.., "date_to":..} برای range یا {"period_name":..} برای period
        """
        report = lambda pct, msg: BackupManager._report(progress_cb, pct, msg)
        report(2, "در حال آماده‌سازی...")

        db = get_session()
        try:
            all_policies = db.query(Policy).all()
            if scope == "range":
                d_from = en_digits(params.get("date_from") or "")
                d_to = en_digits(params.get("date_to") or "")
                policies = [p for p in all_policies if (not d_from or en_digits(p.issue_date or "") >= d_from)
                            and (not d_to or en_digits(p.issue_date or "") <= d_to)]
            elif scope == "period":
                per_name = params.get("period_name")
                policies = [p for p in all_policies if p.period_name == per_name]
            else:
                policies = all_policies
            report(10, "خواندن بیمه‌نامه‌ها...")

            policy_ids = {p.id for p in policies}
            company_ids = {p.company_id for p in policies if p.company_id}
            companies = db.query(Company).filter(Company.id.in_(company_ids)).all() if company_ids else []
            report(18, "خواندن شرکت‌ها...")

            installments = db.query(Installment).filter(Installment.policy_id.in_(policy_ids)).all() if policy_ids else []
            report(28, "خواندن اقساط...")

            receipt_ids = {i.receipt_us_id for i in installments if i.receipt_us_id}
            receipt_ids |= {i.receipt_pasargad_id for i in installments if i.receipt_pasargad_id}
            receipts = db.query(Receipt).filter(Receipt.id.in_(receipt_ids)).all() if receipt_ids else []
            report(36, "خواندن رسیدها...")

            inst_ids = {i.id for i in installments}
            payment_logs = db.query(PaymentLog).filter(PaymentLog.installment_id.in_(inst_ids)).all() if inst_ids else []
            report(44, "خواندن تاریخچه پرداخت‌ها...")

            period_names = {p.period_name for p in policies if p.period_name}
            invoices = db.query(Invoice).filter(Invoice.period_name.in_(period_names)).all() if period_names else []
            report(50, "خواندن صورتحساب‌ها...")

            data = {
                "export_version": 1,
                "scope": scope,
                "params": params,
                "companies": [BackupManager._ser_company(c) for c in companies],
                "policies": [BackupManager._ser_policy(p) for p in policies],
                "installments": [BackupManager._ser_installment(i) for i in installments],
                "receipts": [BackupManager._ser_receipt(r) for r in receipts],
                "payment_logs": [BackupManager._ser_payment_log(lg) for lg in payment_logs],
                "invoices": [BackupManager._ser_invoice(inv) for inv in invoices],
            }
            if scope == "all":
                data["settings"] = load_settings()
        finally:
            db.close()

        report(58, "نوشتن فایل داده...")
        tmp_dir = tempfile.mkdtemp(prefix="mammut_export_")
        try:
            json_path = os.path.join(tmp_dir, "data.json")
            with open(json_path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)

            report(65, "کپی بایگانی و فایل‌ها...")
            with zipfile.ZipFile(dest_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
                zf.write(json_path, arcname="data.json")

                if scope == "all":
                    if os.path.isdir(DATA_ROOT):
                        for root, dirs, files in os.walk(DATA_ROOT):
                            for fn in files:
                                if fn.startswith("~tmp_"): continue
                                full = os.path.join(root, fn)
                                rel = os.path.relpath(full, DATA_ROOT)
                                zf.write(full, arcname=os.path.join("files", rel))
                else:
                    archive_root = os.path.join(DATA_ROOT, BackupManager.ARCHIVE_FOLDER)
                    included_years = set()
                    for per in period_names:
                        year = BackupManager._period_year(per)
                        if not year: continue
                        included_years.add(year)
                        per_dir = os.path.join(archive_root, year, per)
                        if os.path.isdir(per_dir):
                            for root, dirs, files in os.walk(per_dir):
                                for fn in files:
                                    if fn.startswith("~tmp_"): continue
                                    full = os.path.join(root, fn)
                                    rel = os.path.relpath(full, DATA_ROOT)
                                    zf.write(full, arcname=os.path.join("files", rel))
                    for extra in BackupManager.ALWAYS_INCLUDE_FOLDERS:
                        extra_dir = os.path.join(DATA_ROOT, extra)
                        if os.path.isdir(extra_dir):
                            for root, dirs, files in os.walk(extra_dir):
                                for fn in files:
                                    if fn.startswith("~tmp_"): continue
                                    full = os.path.join(root, fn)
                                    rel = os.path.relpath(full, DATA_ROOT)
                                    zf.write(full, arcname=os.path.join("files", rel))
            report(97, "نهایی‌سازی...")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

        report(100, "پایان یافت")
        return {
            "success": True,
            "companies": len(data["companies"]), "policies": len(data["policies"]),
            "installments": len(data["installments"]), "invoices": len(data["invoices"]),
        }

    # ------------------------------- Import -------------------------------
    @staticmethod
    def import_backup(zip_path, progress_cb=None):
        report = lambda pct, msg: BackupManager._report(progress_cb, pct, msg)
        report(2, "در حال استخراج فایل پشتیبان...")

        if not os.path.exists(zip_path):
            return {"success": False, "error": "فایل انتخاب شده یافت نشد."}

        tmp_dir = tempfile.mkdtemp(prefix="mammut_import_")
        created_policies = 0
        created_installments = 0
        try:
            try:
                with zipfile.ZipFile(zip_path, "r") as zf:
                    zf.extractall(tmp_dir)
            except zipfile.BadZipFile:
                return {"success": False, "error": "فایل انتخاب شده یک فایل پشتیبان معتبر (ZIP) نیست."}

            report(12, "خواندن داده‌ها...")
            json_path = os.path.join(tmp_dir, "data.json")
            if not os.path.exists(json_path):
                return {"success": False, "error": "ساختار فایل پشتیبان نامعتبر است (data.json یافت نشد)."}

            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            db = get_session()
            try:
                company_map = {}
                for c in data.get("companies", []):
                    existing = db.query(Company).filter(Company.name == c.get("name")).first()
                    if existing:
                        company_map[c["id"]] = existing.id
                    else:
                        new_c = Company(name=c.get("name"), code=c.get("code"), phone=c.get("phone"))
                        db.add(new_c); db.flush()
                        company_map[c["id"]] = new_c.id
                report(25, "بازیابی شرکت‌ها...")

                policy_map = {}
                for p in data.get("policies", []):
                    existing = None
                    if p.get("policy_number"):
                        existing = db.query(Policy).filter(Policy.policy_number == p["policy_number"]).first()
                    if existing:
                        policy_map[p["id"]] = existing.id
                        continue
                    new_p = Policy(
                        company_id=company_map.get(p.get("company_id")),
                        source_batch_id=p.get("source_batch_id"),
                        policy_number=p.get("policy_number"), contract_number=p.get("contract_number"),
                        insured_name=p.get("insured_name"), premium_amount=p.get("premium_amount") or 0,
                        issue_date=p.get("issue_date"), national_code=p.get("national_code"),
                        plate=p.get("plate"), personnel_code=p.get("personnel_code"),
                        policy_type=p.get("policy_type"), contract_company=p.get("contract_company"),
                        period_name=p.get("period_name"), is_invoiced=bool(p.get("is_invoiced")),
                        invoice_number=p.get("invoice_number"),
                    )
                    db.add(new_p); db.flush()
                    policy_map[p["id"]] = new_p.id
                    created_policies += 1
                report(45, "بازیابی بیمه‌نامه‌ها...")

                receipt_map = {}
                for r in data.get("receipts", []):
                    new_r = Receipt(
                        receipt_type=r.get("receipt_type"), amount=r.get("amount"),
                        payment_date=r.get("payment_date"), description=r.get("description"),
                        file_path=r.get("file_path"), target=r.get("target"),
                    )
                    db.add(new_r); db.flush()
                    receipt_map[r["id"]] = new_r.id
                report(55, "بازیابی رسیدها...")

                inst_map = {}
                for i in data.get("installments", []):
                    new_pol_id = policy_map.get(i.get("policy_id"))
                    if not new_pol_id: continue
                    existing = None
                    if i.get("tracking_id"):
                        existing = db.query(Installment).filter(Installment.tracking_id == i["tracking_id"]).first()
                    if existing:
                        inst_map[i["id"]] = existing.id
                        continue
                    new_i = Installment(
                        tracking_id=i.get("tracking_id"), policy_id=new_pol_id,
                        installment_number=i.get("installment_number"), amount=i.get("amount"),
                        paid_amount_us=i.get("paid_amount_us") or 0, paid_amount_pasargad=i.get("paid_amount_pasargad") or 0,
                        due_date=i.get("due_date"), status_us=i.get("status_us") or "open",
                        status_pasargad=i.get("status_pasargad") or "open",
                        receipt_us_id=receipt_map.get(i.get("receipt_us_id")),
                        receipt_pasargad_id=receipt_map.get(i.get("receipt_pasargad_id")),
                    )
                    db.add(new_i); db.flush()
                    inst_map[i["id"]] = new_i.id
                    created_installments += 1
                report(70, "بازیابی اقساط...")

                for lg in data.get("payment_logs", []):
                    new_inst_id = inst_map.get(lg.get("installment_id"))
                    if not new_inst_id: continue
                    db.add(PaymentLog(
                        installment_id=new_inst_id, receipt_id=receipt_map.get(lg.get("receipt_id")),
                        target=lg.get("target"), amount=lg.get("amount"),
                        registered_by=lg.get("registered_by"),
                    ))
                report(80, "بازیابی تاریخچه پرداخت‌ها...")

                for inv in data.get("invoices", []):
                    existing = None
                    if inv.get("invoice_number"):
                        existing = db.query(Invoice).filter(Invoice.invoice_number == inv["invoice_number"]).first()
                    if existing: continue
                    db.add(Invoice(
                        company_name=inv.get("company_name"), invoice_number=inv.get("invoice_number"),
                        period_name=inv.get("period_name"), total_policies=inv.get("total_policies") or 0,
                        total_amount=inv.get("total_amount") or 0, created_date=inv.get("created_date"),
                        folder_path=inv.get("folder_path"), status=inv.get("status") or "open",
                    ))
                report(88, "بازیابی صورتحساب‌ها...")

                db.commit()
            except Exception:
                db.rollback()
                raise
            finally:
                db.close()

            if "settings" in data:
                try:
                    s = load_settings()
                    s.update(data["settings"])
                    save_settings(s)
                except Exception:
                    pass

            report(93, "بازیابی بایگانی و فایل‌ها...")
            files_dir = os.path.join(tmp_dir, "files")
            if os.path.isdir(files_dir):
                for root, dirs, files in os.walk(files_dir):
                    rel_root = os.path.relpath(root, files_dir)
                    target_root = DATA_ROOT if rel_root == "." else os.path.join(DATA_ROOT, rel_root)
                    os.makedirs(target_root, exist_ok=True)
                    for fn in files:
                        src = os.path.join(root, fn)
                        dst = os.path.join(target_root, fn)
                        try:
                            shutil.copy2(src, dst)
                        except Exception:
                            pass
            report(99, "پاکسازی...")
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

        report(100, "پایان یافت")
        return {"success": True, "policies_created": created_policies, "installments_created": created_installments}
